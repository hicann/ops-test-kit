#!/usr/bin/env python3
# -*- coding: UTF-8 -*-
# Copyright (c) 2026 Huawei Technologies Co., Ltd.
# This program is free software, you can redistribute it and/or modify it under the terms and conditions of
# CANN Open Software License Agreement Version 2.0 (the "License").
# Please refer to the License for details. You may not use this file except in compliance with the License.
# THIS SOFTWARE IS PROVIDED ON AN "AS IS" BASIS, WITHOUT WARRANTIES OF ANY KIND, EITHER EXPRESS OR IMPLIED,
# INCLUDING BUT NOT LIMITED TO NON-INFRINGEMENT, MERCHANTABILITY, OR FITNESS FOR A PARTICULAR PURPOSE.
# See LICENSE in the root of the software repository for the full text of the License.
"""ttk list 重写为 raw 读：只读 testcase_name 列，其他列不解析不校验。"""

import csv
import subprocess
import sys

HDR = ["testcase_name", "op_name", "is_enabled", "input_shapes", "input_dtypes"]


def _write_csv(path, rows):
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(HDR)
        w.writerows(rows)


def _run_list(*args):
    return subprocess.run(
        [sys.executable, "-m", "ttk", "list", *args],
        capture_output=True,
        text=True,
        cwd=str(__import__("pathlib").Path(__file__).parents[2]),
        check=False,
    )


def test_list_csv_names_and_total(tmp_path):
    f = tmp_path / "c.csv"
    _write_csv(
        f, [["a1", "add", "true", "((2,3),)", "('float32',)"], ["a2", "sub", "true", "((2,3),)", "('float32',)"]]
    )
    r = _run_list("-i", str(f))
    assert r.returncode == 0
    assert "a1\n" in r.stdout

    assert "a2\n" in r.stdout
    assert "Total: 2 case(s)" in r.stdout


def test_list_dirty_columns_ignored(tmp_path):
    f = tmp_path / "c.csv"
    _write_csv(f, [["a1", "add", "true", "NOT_A_SHAPE", "BAD_DTYPE"]])
    r = _run_list("-i", str(f))
    assert r.returncode == 0  # 其他列不解析，脏数据不影响
    assert "a1" in r.stdout


def test_list_missing_testcase_name_column_errors(tmp_path):
    f = tmp_path / "c.csv"
    with open(f, "w", newline="") as fh:
        csv.writer(fh).writerows([["op_name"], ["add"]])
    r = _run_list("-i", str(f))
    assert r.returncode != 0
    assert "testcase_name" in r.stderr


def test_list_empty_name_cell_errors(tmp_path):
    f = tmp_path / "c.csv"
    _write_csv(f, [["", "add", "true", "((2,3),)", "('float32',)"]])
    r = _run_list("-i", str(f))
    assert r.returncode != 0
    assert "testcase_name" in r.stderr


def test_list_filters_disabled_rows(tmp_path):
    f = tmp_path / "c.csv"
    _write_csv(
        f, [["a1", "add", "true", "((2,3),)", "('float32',)"], ["a2", "sub", "false", "((2,3),)", "('float32',)"]]
    )
    r = _run_list("-i", str(f))
    assert r.returncode == 0
    assert "a1" in r.stdout

    assert "a2" not in r.stdout
    assert "Total: 1 case(s)" in r.stdout


def test_list_op_filter_and_comma_multi(tmp_path):
    f = tmp_path / "c.csv"
    _write_csv(
        f,
        [
            ["a1", "add", "true", "((2,3),)", "('float32',)"],
            ["a2", "sub", "true", "((2,3),)", "('float32',)"],
            ["a3", "mul", "true", "((2,3),)", "('float32',)"],
        ],
    )
    r = _run_list("-i", str(f), "--op", "add,sub")
    assert r.returncode == 0
    assert "a1" in r.stdout

    assert "a2" in r.stdout

    assert "a3" not in r.stdout


def test_list_op_filter_excludes_disabled_rows(tmp_path):
    # --op 与 is_enabled 叠加：禁用行即使 op 匹配也不出现（保持旧行为, testcase_manager.py:469 无条件过滤）
    f = tmp_path / "c.csv"
    _write_csv(
        f, [["a1", "add", "true", "((2,3),)", "('float32',)"], ["a2", "add", "false", "((2,3),)", "('float32',)"]]
    )
    r = _run_list("-i", str(f), "--op", "add")
    assert r.returncode == 0
    assert "a1" in r.stdout

    assert "a2" not in r.stdout


def test_list_op_filter_api_name_column(tmp_path):
    f = tmp_path / "c.csv"
    with open(f, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["testcase_name", "api_name"])
        w.writerow(["x1", "aclnnAdd"])
        w.writerow(["x2", "aclnnSub"])
    r = _run_list("-i", str(f), "--op", "aclnnAdd")
    assert r.returncode == 0
    assert "x1" in r.stdout

    assert "x2" not in r.stdout


def test_list_op_filter_prefers_op_name_column(tmp_path):
    # 表头同时含 op_name 与 api_name 两列时，--op 按 op_name 列过滤（op_name 列优先）
    f = tmp_path / "c.csv"
    with open(f, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["testcase_name", "op_name", "api_name"])
        w.writerow(["p1", "add", "aclnnAdd"])
    r = _run_list("-i", str(f), "--op", "add")
    assert r.returncode == 0
    assert "Total: 1 case(s)" in r.stdout
    r = _run_list("-i", str(f), "--op", "aclnnAdd")
    assert r.returncode == 0
    assert "Total: 0 case(s)" in r.stdout


def test_list_op_filter_no_column_zero_match(tmp_path):
    f = tmp_path / "c.csv"
    with open(f, "w", newline="") as fh:
        csv.writer(fh).writerows([["testcase_name"], ["a1"]])
    r = _run_list("-i", str(f), "--op", "add")
    assert r.returncode == 0
    assert "Total: 0 case(s)" in r.stdout


def test_list_duplicate_names_unchanged(tmp_path):
    f = tmp_path / "c.csv"
    _write_csv(
        f, [["dup", "add", "true", "((2,3),)", "('float32',)"], ["dup", "sub", "true", "((2,3),)", "('float32',)"]]
    )
    r = _run_list("-i", str(f))
    assert r.returncode == 0
    assert r.stdout.count("dup") == 2


def test_list_header_only_csv_zero_total(tmp_path):
    f = tmp_path / "c.csv"
    with open(f, "w", newline="") as fh:
        csv.writer(fh).writerow(HDR)
    r = _run_list("-i", str(f))
    assert r.returncode == 0
    assert "Total: 0 case(s)" in r.stdout


def test_list_xlsx_sheet_default_and_named(tmp_path):
    import openpyxl

    f = tmp_path / "c.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.append(["testcase_name", "op_name"])
    ws1.append(["s1_case", "add"])
    ws2 = wb.create_sheet("T2")
    ws2.append(["testcase_name", "op_name"])
    ws2.append(["s2_case", "sub"])
    wb.save(f)

    r = _run_list("-i", str(f))
    assert r.returncode == 0
    assert "s1_case" in r.stdout  # 默认首表

    assert "s2_case" not in r.stdout
    r = _run_list("-i", str(f), "--sheet", "T2")
    assert r.returncode == 0
    assert "s2_case" in r.stdout

    assert "s1_case" not in r.stdout


def test_list_xlsx_sheet_not_found_errors(tmp_path):
    import openpyxl

    f = tmp_path / "c.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["testcase_name"])
    wb.active.append(["a1"])
    wb.save(f)
    r = _run_list("-i", str(f), "--sheet", "NOPE")
    assert r.returncode != 0


def test_list_xlsm_supported(tmp_path):
    import openpyxl

    f = tmp_path / "c.xlsm"
    wb = openpyxl.Workbook()
    wb.active.append(["testcase_name", "op_name"])
    wb.active.append(["m1", "add"])
    wb.save(f)  # openpyxl 按后缀保存, 全新 Workbook 存 .xlsm 无需额外参数
    r = _run_list("-i", str(f))
    assert r.returncode == 0
    assert "m1" in r.stdout


def test_list_is_enabled_invalid_values_treated_enabled(tmp_path):
    # 非法值宽容化: 仅 "false"/"0" 视为禁用, 其余（旧路径崩溃或 eval 语义）一律视为启用
    f = tmp_path / "c.csv"
    _write_csv(
        f,
        [
            ["a1", "add", "no", "((2,3),)", "('float32',)"],
            ["a2", "sub", "off", "((2,3),)", "('float32',)"],
            ["a3", "mul", "1.5", "((2,3),)", "('float32',)"],
        ],
    )
    r = _run_list("-i", str(f))
    assert r.returncode == 0
    assert "a1" in r.stdout

    assert "a2" in r.stdout

    assert "a3" in r.stdout
    assert "Total: 3 case(s)" in r.stdout


def test_list_empty_table_errors(tmp_path):
    # 空文件（无任何行）: read_table raise ValueError（区别于仅表头）
    f = tmp_path / "c.csv"
    f.write_text("")
    r = _run_list("-i", str(f))
    assert r.returncode != 0
